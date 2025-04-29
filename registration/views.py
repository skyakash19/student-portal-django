import datetime
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth import login, update_session_auth_hash
from django.contrib.auth.decorators import login_required
from django.contrib.admin.views.decorators import staff_member_required
from django.contrib import messages
from django.http import HttpResponse
from django.template.loader import get_template
from django.db.models import Count
from django.contrib.auth.models import User
from django.core.paginator import Paginator
from django.db.models import Q
from django.core.files.storage import FileSystemStorage
import pandas as pd
from xhtml2pdf import pisa  # type: ignore
from datetime import datetime
from django.utils import timezone
from django.utils.timezone import make_aware
from .models import Course, Student, UnenrollRequest
from .forms import ProfileForm, StudentForm, CustomSignupForm
from .models import Course, Enrollment
from django.contrib.auth import get_user_model
from django.contrib.auth.decorators import login_required

# -------------------------------
# Home View
# -------------------------------
def home(request):
    if request.user.is_authenticated:
        return redirect('dashboard')
    return render(request, 'home.html')


# -------------------------------
# Dashboard
# -------------------------------
@login_required
def dashboard(request):
    user = request.user
    context = {'user': user}

    if not user.is_superuser:
        try:
            student = Student.objects.get(user=user)
            context['student'] = student
            context['profile'] = student  # Add profile context for template use
        except Student.DoesNotExist:
            context['student'] = None
            context['profile'] = None
    else:
        context['profile'] = None  # For superuser, no profile photo

    return render(request, 'dashboard.html', context)

# -------------------------------
# Sign Up
# -------------------------------
def signup_view(request):
    if request.method == 'POST':
        form = CustomSignupForm(request.POST, request.FILES)
        if form.is_valid():
            user = form.save(commit=False)
            user.email = form.cleaned_data['email']
            user.save()

            bio = form.cleaned_data.get('bio')
            profile_photo = form.cleaned_data.get('profile_photo')
            Student.objects.create(user=user, bio=bio, profile_photo=profile_photo)

            messages.success(request, "✅ Account created successfully! Please log in.")
            return redirect('login')
    else:
        form = CustomSignupForm()
    return render(request, 'signup.html', {'form': form})


# -------------------------------
# Course Stats
# -------------------------------
@login_required
def course_stats(request):
    total_courses = Course.objects.count()
    total_students = Student.objects.count()

    course_list = Course.objects.annotate(student_count=Count('student'))

    return render(request, 'course_stats.html', {
        'total_courses': total_courses,
        'total_students': total_students,
        'course_list': course_list,
    })


# -------------------------------
# Browse Courses
# -------------------------------
@login_required
def browse_courses(request):
    courses = Course.objects.all()
    try:
        student = Student.objects.get(user=request.user)
        enrolled_courses = student.courses.all()
    except Student.DoesNotExist:
        enrolled_courses = []
    return render(request, 'browse_courses.html', {
        'courses': courses,
        'enrolled_courses': enrolled_courses,
    })


# -------------------------------
# My Courses
# -------------------------------
@login_required
def my_courses(request):
    try:
        student = Student.objects.select_related('user').prefetch_related('courses').get(user=request.user)
        courses = student.courses.all()
        
        if not courses:
            messages.info(request, "You're not enrolled in any courses yet.")
        
        return render(request, 'my_courses.html', {'courses': courses})
    
    except Student.DoesNotExist:
        messages.error(request, "Student profile not found. Please contact support.")
        return redirect('dashboard')  # or wherever you want to redirect

# -------------------------------
# Enroll in Course
# -------------------------------
@login_required
def enroll_course(request, course_id):
    try:
        course = Course.objects.get(id=course_id)
        student = Student.objects.get(user=request.user)
        student.courses.add(course)
        return redirect('my_courses')
    except (Course.DoesNotExist, Student.DoesNotExist):
        return HttpResponse("Something went wrong.")


# -------------------------------
# Profile
# -------------------------------
@login_required
def profile(request):
    user = request.user
    is_superuser = user.is_superuser
    student = None

    if not is_superuser:
        try:
            student = Student.objects.get(user=user)
        except Student.DoesNotExist:
            messages.warning(request, "⚠️ Student profile not found.")

    if request.method == 'POST':
        username = request.POST.get('username', '').strip()
        first_name = request.POST.get('first_name', '').strip()
        last_name = request.POST.get('last_name', '').strip()
        email = request.POST.get('email', '').strip()
        password1 = request.POST.get('password1', '')
        password2 = request.POST.get('password2', '')

        if not username or not email:
            messages.error(request, "❌ Username and email are required.")
            return redirect('profile')

        user.username = username
        user.first_name = first_name
        user.last_name = last_name
        user.email = email

        if password1 or password2:
            if password1 == password2:
                user.set_password(password1)
                update_session_auth_hash(request, user)
                messages.success(request, "🔐 Password updated successfully.")
            else:
                messages.error(request, "❌ Passwords do not match.")
                return redirect('profile')

        user.save()

        if student:
            profile_photo = request.FILES.get('image')
            if profile_photo:
                student.profile_photo = profile_photo
            student.save()

        messages.success(request, "✅ Profile updated successfully.")
        return redirect('profile')

    enrolled_courses = student.courses.all() if student else []
    return render(request, 'profile.html', {
        'user': user,
        'profile': student,
        'enrolled_courses': enrolled_courses,
        'is_superuser': is_superuser,
    })


# -------------------------------
# Admin – Course Management
# -------------------------------
@staff_member_required
def add_course(request):
    if request.method == 'POST':
        name = request.POST.get('name')
        description = request.POST.get('description')
        if name:
            Course.objects.create(name=name, description=description)
            messages.success(request, "✅ Course added successfully!")
            return redirect('add_course')
    courses = Course.objects.all()
    return render(request, 'add_course.html', {'courses': courses})

@staff_member_required
def edit_student(request, student_id):
    student = get_object_or_404(Student, id=student_id)
    user = student.user  # Assuming Student has OneToOneField to User

    if request.method == 'POST':
        username = request.POST.get('username')
        email = request.POST.get('email')

        if username and email:
            user.username = username
            user.email = email
            user.save()

        return redirect('student_list')  # After saving, redirect

    return redirect('student_list')  # fallback



@staff_member_required
def delete_student(request, student_id):
    student = get_object_or_404(Student, id=student_id)
    if request.method == 'POST':  # It's good practice to use POST for deletion
        student.delete()
        # Optionally, add a success message
        return redirect('student_list')  # Redirect back to the student list
    # If it's a GET request, you might want to show a confirmation page
    return render(request, 'delete_confirm.html', {'student': student})

@staff_member_required
def edit_course(request, course_id):
    course = get_object_or_404(Course, id=course_id)
    if request.method == 'POST':
        course.name = request.POST.get('name')
        course.description = request.POST.get('description')
        course.save()
        messages.success(request, '✅ Course updated successfully!')
        return redirect('add_course')
    return render(request, 'edit_course.html', {'course': course})


@staff_member_required
def delete_course(request, course_id):
    course = get_object_or_404(Course, id=course_id)
    course.delete()
    messages.success(request, '🗑️ Course deleted successfully!')
    return redirect('add_course')


# -------------------------------
# Admin – Student List
# -------------------------------
from django.contrib.admin.views.decorators import staff_member_required
from django.core.paginator import Paginator
from django.db.models import Q
from django.shortcuts import render
from .models import Student  # make sure Student model is imported

@staff_member_required
def student_list(request):
    query = request.GET.get('q', '')
    students = Student.objects.select_related('user')

    if query:
        students = students.filter(
            Q(user__username__icontains=query) |
            Q(user__email__icontains=query)
        )

    paginator = Paginator(students, 10)  # ✅ FIXED
    page_number = request.GET.get('page')
    page_obj = paginator.get_page(page_number)

    return render(request, 'student_list.html', {
        'students': page_obj,
        'query': query
    })

# -------------------------------
# Export Students to Excel
# -------------------------------
import pandas as pd
from django.http import HttpResponse
from django.contrib.auth.decorators import login_required
from django.utils.timezone import localtime
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from .models import Student


@login_required
def export_students_excel(request):
    # Fetch all students with related user data
    students = Student.objects.select_related('user').all()

    # Prepare structured data
    data = [
        {
            'Username': student.user.username,
            'Email': student.user.email,
            'Date Joined': localtime(student.user.date_joined).strftime('%Y-%m-%d %H:%M:%S')
        }
        for student in students
    ]

    df = pd.DataFrame(data)

    # Set up HTTP response
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename=Registered_Students_Export.xlsx'

    # Write Excel using openpyxl for styling
    with pd.ExcelWriter(response, engine='openpyxl') as writer:
        df.to_excel(writer, index=False, sheet_name='Students')
        workbook = writer.book
        worksheet = writer.sheets['Students']

        # Style header row
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")

        for col_num, column_title in enumerate(df.columns, 1):
            cell = worksheet.cell(row=1, column=col_num)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")

            # Adjust column width dynamically
            max_length = max(df[column_title].astype(str).map(len).max(), len(column_title))
            adjusted_width = max_length + 4
            column_letter = get_column_letter(col_num)
            worksheet.column_dimensions[column_letter].width = adjusted_width

        # Apply zebra striping (optional flair)
        for row in worksheet.iter_rows(min_row=2, max_row=worksheet.max_row):
            for cell in row:
                cell.alignment = Alignment(horizontal="left")
            if row[0].row % 2 == 0:
                for cell in row:
                    cell.fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")

    return response



# -------------------------------
# Export Students to PDF
# -------------------------------
from django.utils.timezone import localtime
from io import BytesIO
from xhtml2pdf import pisa

@staff_member_required
def export_students_pdf(request):
    students = Student.objects.select_related('user').all().order_by('-user__date_joined')

    context = {
        'students': students,
        'generated_on': localtime().strftime('%B %d, %Y • %I:%M %p'),
        'admin': request.user.get_full_name() or request.user.username
    }

    template = get_template('export_students_pdf.html')
    html = template.render(context)

    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = 'attachment; filename="Registered_Students_Dashboard.pdf"'

    pisa_status = pisa.CreatePDF(BytesIO(html.encode('UTF-8')), dest=response, encoding='UTF-8')
    if pisa_status.err:
        return HttpResponse('❌ Error generating PDF. Please try again.', status=500)
    return response

from django.contrib.admin.views.decorators import staff_member_required
from django.template.loader import get_template
from django.http import HttpResponse
from django.utils.timezone import localtime
from xhtml2pdf import pisa
from io import BytesIO
from django.shortcuts import get_object_or_404
from .models import Student

@staff_member_required
def export_single_student_pdf(request, student_id):
    student = get_object_or_404(Student, pk=student_id)

    context = {
        'student': student,
        'generated_on': localtime().strftime('%B %d, %Y • %I:%M %p'),
        'admin': request.user.get_full_name() or request.user.username
    }

    template = get_template('export_single_student_pdf.html')
    html = template.render(context)

    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="{student.user.username}_profile.pdf"'

    pisa_status = pisa.CreatePDF(BytesIO(html.encode('UTF-8')), dest=response, encoding='UTF-8')
    if pisa_status.err:
        return HttpResponse('❌ Error generating PDF. Please try again.', status=500)
    return response

from openpyxl import Workbook
from django.http import HttpResponse
from django.shortcuts import get_object_or_404
from .models import Student
from django.contrib.admin.views.decorators import staff_member_required

@staff_member_required
def export_single_student_excel(request, student_id):
    student = get_object_or_404(Student, pk=student_id)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Student Profile"

    # Header
    sheet.append(["Field", "Value"])

    # Student data
    sheet.append(["Username", student.user.username])
    sheet.append(["Email", student.user.email])
    sheet.append(["Full Name", student.user.get_full_name()])
    sheet.append(["Date Joined", student.user.date_joined.strftime('%B %d, %Y')])

    # Return Excel file
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename={student.user.username}_profile.xlsx'
    workbook.save(response)
    return response


from django.shortcuts import render, redirect
from django.contrib import messages
from .models import Enrollment, Course

def unenroll_course(request, course_id):
    try:
        # Fetch the course object using the course_id
        course = Course.objects.get(id=course_id)
        
        # Check if the user has a related student object
        if not hasattr(request.user, 'student'):
            messages.error(request, "You must be logged in as a student to unenroll from courses.")
            return redirect('my_courses')
        
        # Fetch the student instance
        student = request.user.student

        # Check if the student is enrolled in the course
        enrollment = Enrollment.objects.get(student=student, course=course)
        
        # Unenroll logic: delete the enrollment
        enrollment.delete()

        messages.success(request, f"You have been successfully unenrolled from {course.name}.")
        return redirect('my_courses')  # Redirect back to the user's courses page

    except Enrollment.DoesNotExist:
        messages.error(request, "You are not enrolled in this course.")
        return redirect('my_courses')

    except Course.DoesNotExist:
        messages.error(request, "Course not found.")
        return redirect('my_courses')

    except Exception as e:
        messages.error(request, f"An unexpected error occurred: {e}")
        return redirect('my_courses')

from django.contrib import messages
from django.shortcuts import redirect, get_object_or_404


@login_required
def request_unenroll_course(request, course_id):
    course = get_object_or_404(Course, id=course_id)
    
    # Check if already requested
    if UnenrollRequest.objects.filter(user=request.user, course=course).exists():
        messages.warning(request, "You have already requested to unenroll from this course.")
    else:
        UnenrollRequest.objects.create(user=request.user, course=course)
        messages.success(request, f"Unenrollment request for '{course.name}' sent successfully.")
    
    return redirect('my_courses')


from django.contrib.admin.views.decorators import staff_member_required

@staff_member_required
def manage_unenroll_requests(request):
    requests = UnenrollRequest.objects.filter(is_resolved=False).select_related('user', 'course')
    return render(request, 'admin_unenroll_requests.html', {'requests': requests})

@staff_member_required
def approve_unenroll_request(request, req_id):
    req = get_object_or_404(UnenrollRequest, id=req_id)
    req.is_resolved = True
    req.is_approved = True
    req.save()
    req.course.students.remove(req.user)  # assuming `students` is a ManyToMany field
    messages.success(request, f"{req.user.username} has been unenrolled from {req.course.name}.")
    return redirect('manage_unenroll_requests')

@staff_member_required
def deny_unenroll_request(request, req_id):
    req = get_object_or_404(UnenrollRequest, id=req_id)
    req.is_resolved = True
    req.is_approved = False
    req.save()
    messages.warning(request, f"Request to unenroll {req.user.username} from {req.course.name} was denied.")
    return redirect('manage_unenroll_requests')
